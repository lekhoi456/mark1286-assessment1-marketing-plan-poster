"""Extract Assessment 1 MARK1286 teaching materials into Markdown.

Run from the workspace root (assessment1-marketing-plan-poster):

    uv run --with python-pptx python -B 06_workflow/scripts/extract_materials.py [--force] [--slug SLUG]

Reads ../study-hub/content/materials.json (the catalogue) and extracts
Assessment 1 entries from the ORIGINAL file in the course root (CR):

- .pptx: per slide - number, title, body text (text frames, tables, grouped
  shapes and SmartArt, in reading order), hyperlinks, image alt text and
  speaker notes.
- .pdf: per page via `pdftotext -layout`. A page whose text layer has fewer
  than OCR_BELOW words (for example `w04-tesla`, whose body text is vector
  outlines) is rendered with `pdftoppm` into a temporary folder outside the
  workspace and read with `tesseract`. If OCR finds clearly more words than
  the text layer, the page is marked "(OCR)" and the OCR text is written
  first, followed by the text layer; otherwise the text layer is kept.
- .doc/.docx: `textutil -convert txt`.
- .jpg/.jpeg/.png: a stub only (image path and pixel size). The Assessment 1
  sample posters are photographs of hand-drawn posters; their handwritten
  text is not machine-extracted (read the image; see
  00_brief_and_criteria/exemplars.md).

Output: 03_course_materials/extracted/<slug>.md, each with a header (slug,
title, original and study-hub paths, extraction method, slide or page count,
OCR pages). Non-breaking spaces are written as ordinary spaces so that quotes
can be grepped; all other characters are kept as extracted. Idempotent: an
output that is newer than both its source file and this script is skipped
unless --force is given. Repeat --slug to restrict the run; positional slugs
also work. Weeks 7–9 and Assessment 2 vlogs are excluded, including explicit
selection. Exit status 1 on extraction failure; invalid selection exits 2.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

WS = Path(__file__).resolve().parents[2]
CR = WS.parent
CATALOGUE = CR / "study-hub" / "content" / "materials.json"
HUB_PDF_DIR = CR / "study-hub" / "public" / "materials"
OUT_DIR = WS / "03_course_materials" / "extracted"
SCRIPT = Path(__file__).resolve()

OCR_BELOW = 100  # text-layer words per page below which a page is also OCR'd
OCR_DPI = 300
OCR_LANG = "eng"
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}

DGM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def clean(text: str) -> str:
    """Normalise in-paragraph line breaks and trailing space."""
    return text.replace("\v", " ").replace("\r", " ").strip()


def pdf_page_count(path: Path) -> int | None:
    if not path.exists():
        return None
    out = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True).stdout
    for line in out.splitlines():
        if line.startswith("Pages:"):
            return int(line.split()[1])
    return None


def word_count(text: str) -> int:
    return sum(1 for token in text.split() if any(c.isalnum() for c in token))


# --------------------------------------------------------------------------- pptx


def paragraph_lines(text_frame, indent: str = "") -> list[str]:
    lines = []
    for para in text_frame.paragraphs:
        text = clean("".join(run.text for run in para.runs) or para.text)
        if not text:
            continue
        links = []
        for run in para.runs:
            try:
                address = run.hyperlink.address
            except Exception:  # noqa: BLE001 - malformed rels in some decks
                address = None
            if address and address not in links:
                links.append(address)
        suffix = "".join(f" <{link}>" for link in links)
        lines.append(f"{indent}{'  ' * para.level}- {text}{suffix}")
    return lines


def smartart_lines(shape, part) -> list[str]:
    """Return the text nodes of a SmartArt graphic from its data part."""
    lines = []
    for rel_ids in shape._element.iter(f"{{{NS_DGM}}}relIds"):
        dm = rel_ids.get(f"{{{NS_R}}}dm")
        if not dm or dm not in part.rels:
            continue
        from lxml import etree

        root = etree.fromstring(part.rels[dm].target_part.blob)
        for pt in root.iter(f"{{{NS_DGM}}}pt"):
            if pt.get("type") not in (None, "node"):
                continue
            for para in pt.iter(f"{{{NS_A}}}p"):
                text = clean("".join(t.text or "" for t in para.iter(f"{{{NS_A}}}t")))
                if text:
                    lines.append(f"- {text}")
    return lines


def table_lines(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            parts = [clean(p.text) for p in cell.text_frame.paragraphs]
            cells.append(" / ".join(p for p in parts if p).replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return out


def sort_key(shape):
    top = shape.top if shape.top is not None else 0
    left = shape.left if shape.left is not None else 0
    return (top, left)


def shape_lines(shapes, part, skip_id=None) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines: list[str] = []
    for shape in sorted(shapes, key=sort_key):
        if skip_id is not None and shape.shape_id == skip_id:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines += shape_lines(shape.shapes, part)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            lines += [""] + table_lines(shape.table) + [""]
            continue
        graphic = shape._element.find(f".//{{{NS_A}}}graphicData")
        if graphic is not None and graphic.get("uri") == DGM_URI:
            lines += smartart_lines(shape, part)
            continue
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = shape.chart
            if chart.has_title and chart.chart_title.has_text_frame:
                lines.append(f"- [Chart: {clean(chart.chart_title.text_frame.text)}]")
            continue
        if shape.has_text_frame:
            lines += paragraph_lines(shape.text_frame)
        try:
            action_link = shape.click_action.hyperlink.address
        except Exception:  # noqa: BLE001
            action_link = None
        if action_link:
            lines.append(f"- [Shape link: <{action_link}>]")
        descr = shape._element.xpath("./*[1]/p:cNvPr/@descr")
        if descr:
            alt = " ".join(descr[0].split("Description automatically generated")[0].split())
            if alt:
                kind = "Media" if shape.shape_type == MSO_SHAPE_TYPE.MEDIA else "Image"
                lines.append(f"- [{kind} alt text: {alt}]")
    return lines


def extract_pptx(src: Path) -> tuple[list[str], int, str, list[int]]:
    from pptx import Presentation

    prs = Presentation(str(src))
    body: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        title = " ".join(title.split()) or "(no title)"
        hidden = slide._element.get("show") == "0"
        body.append(f"## Slide {number} — {title}" + (" [hidden slide]" if hidden else ""))
        body.append("")
        lines = shape_lines(slide.shapes, slide.part, skip_id=title_shape.shape_id if title_shape is not None else None)
        body += lines if lines else ["(no body text)"]
        external = []
        for rel in slide.part.rels.values():
            if rel.is_external and rel.target_ref not in external:
                external.append(rel.target_ref)
        if external:
            body += ["", "**External links on slide:**", ""] + [f"- <{link}>" for link in external]
        if slide.has_notes_slide:
            notes = [clean(p.text) for p in slide.notes_slide.notes_text_frame.paragraphs] if slide.notes_slide.notes_text_frame else []
            notes = [n for n in notes if n]
            if notes:
                body += ["", "**Speaker notes:**", ""] + notes
        body.append("")
    return body, len(prs.slides), "python-pptx (slide text, tables, SmartArt, alt text, speaker notes)", []


# --------------------------------------------------------------------------- pdf / doc / image


def tesseract_version() -> str:
    out = subprocess.run(["tesseract", "--version"], capture_output=True, text=True)
    first = (out.stdout or out.stderr).splitlines()
    return first[0].strip() if first else "tesseract"


def ocr_page(src: Path, number: int, tmp: Path) -> str:
    """Render one PDF page at OCR_DPI into tmp and return tesseract's text."""
    stem = tmp / f"page-{number}"
    subprocess.run(["pdftoppm", "-r", str(OCR_DPI), "-f", str(number), "-l", str(number), "-singlefile", "-png",
                    str(src), str(stem)], capture_output=True, check=True)
    result = subprocess.run(["tesseract", f"{stem}.png", "stdout", "-l", OCR_LANG], capture_output=True, text=True,
                            check=True)
    return result.stdout


def extract_pdf(src: Path) -> tuple[list[str], int, str, list[int]]:
    text = subprocess.run(["pdftotext", "-layout", str(src), "-"], capture_output=True, text=True, check=True).stdout
    pages = text.split("\f")
    count = pdf_page_count(src) or len(pages)
    pages = (pages + [""] * count)[:count]
    body: list[str] = []
    ocr_pages: list[int] = []
    tmp = Path(tempfile.mkdtemp(prefix="mark1286-ocr-"))
    try:
        for number, page in enumerate(pages, start=1):
            layer = page.rstrip() or "(no text on page)"
            text_words = word_count(page)
            ocr_text = ""
            if text_words < OCR_BELOW:
                ocr_text = ocr_page(src, number, tmp).rstrip()
            ocr_words = word_count(ocr_text)
            if ocr_text and ocr_words >= max(2 * text_words, text_words + 50):
                ocr_pages.append(number)
                body += [
                    f"## Page {number} (OCR)",
                    "",
                    f"Text layer: {text_words} words; OCR: {ocr_words} words ({tesseract_version()}, `-l {OCR_LANG}`, "
                    f"page rendered at {OCR_DPI} dpi). Most of this page's text is not in the PDF text layer "
                    "(vector outlines or image). The OCR text below is machine-read: check every quoted word "
                    "against the rendered page before using it.",
                    "",
                    "**OCR text:**",
                    "",
                    "```text",
                    ocr_text,
                    "```",
                    "",
                    "**PDF text layer (`pdftotext -layout`):**",
                    "",
                    "```text",
                    layer,
                    "```",
                    "",
                ]
            else:
                body += [f"## Page {number}", "", "```text", layer, "```", ""]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    method = "pdftotext -layout"
    if ocr_pages:
        method += f"; OCR ({tesseract_version()}, {OCR_DPI} dpi) on page(s) {', '.join(map(str, ocr_pages))}"
    return body, count, method, ocr_pages


def extract_doc(src: Path) -> tuple[list[str], None, str, list[int]]:
    text = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(src)], capture_output=True, text=True, check=True).stdout
    return ["## Full text", "", text.rstrip(), ""], None, "textutil -convert txt", []


def image_size(src: Path) -> str:
    out = subprocess.run(["sips", "-g", "pixelWidth", "-g", "pixelHeight", str(src)], capture_output=True, text=True).stdout
    dims = {}
    for line in out.splitlines():
        parts = line.split(":")
        if len(parts) == 2 and parts[0].strip() in ("pixelWidth", "pixelHeight"):
            dims[parts[0].strip()] = parts[1].strip()
    if len(dims) == 2:
        return f"{dims['pixelWidth']} × {dims['pixelHeight']} pixels"
    return "unknown"


def extract_image(src: Path, entry: dict) -> tuple[list[str], int, str, list[int]]:
    body = [
        "## Image (not machine-extracted)",
        "",
        f"- Image file (relative to course root): `{entry['file']}`",
        f"- Image size: {image_size(src)}",
        f"- Study-hub copy: `study-hub/public/materials/{entry['slug']}{src.suffix.lower()}`",
        "",
        "The image is a photograph of a hand-drawn poster. Its content is handwritten and has NOT been machine-extracted "
        "(no OCR: handwriting OCR would be unreliable). Read the image itself; the human reading and the lessons drawn "
        "from the sample posters are in `00_brief_and_criteria/exemplars.md`. Nothing on this sample may be quoted "
        "from this file.",
        "",
    ]
    return body, 1, "stub (image; handwritten content not machine-extracted)", []


# --------------------------------------------------------------------------- main


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="overwrite outputs even when up to date")
    parser.add_argument("--slug", action="append", default=[], help="extract one A1 catalogue slug (repeatable)")
    parser.add_argument("slugs", nargs="*", help="optional A1 catalogue slugs")
    args = parser.parse_args(argv)
    force = args.force
    only = set(args.slug + args.slugs)
    catalogue = json.loads(CATALOGUE.read_text(encoding="utf-8"))
    catalogue = [
        entry for entry in catalogue
        if (entry.get("week") is None or entry["week"] <= 6)
        and not entry["slug"].startswith("sample-vlog-")
        and "Assessment 2" not in entry["file"]
    ]
    invalid = only - {entry["slug"] for entry in catalogue}
    if invalid:
        parser.error("unknown or out-of-scope A1 slug(s): " + ", ".join(sorted(invalid)))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    written = skipped = failed = 0
    for entry in catalogue:
        slug = entry["slug"]
        if only and slug not in only:
            continue
        src = CR / entry["file"]
        out = OUT_DIR / f"{slug}.md"
        if not src.exists():
            print(f"MISSING SOURCE {slug}: {entry['file']}", file=sys.stderr)
            failed += 1
            continue
        newest_input = max(src.stat().st_mtime, SCRIPT.stat().st_mtime)
        if not force and out.exists() and out.stat().st_mtime > newest_input:
            skipped += 1
            continue
        ext = src.suffix.lower()
        try:
            if ext == ".pptx":
                body, count, method, ocr_pages = extract_pptx(src)
                unit = "Slides"
            elif ext == ".pdf":
                body, count, method, ocr_pages = extract_pdf(src)
                unit = "Pages"
            elif ext in (".doc", ".docx"):
                body, count, method, ocr_pages = extract_doc(src)
                unit = "Pages"
            elif ext in IMAGE_EXTS:
                body, count, method, ocr_pages = extract_image(src, entry)
                unit = "Images"
            else:
                print(f"UNSUPPORTED {slug}: {ext}", file=sys.stderr)
                failed += 1
                continue
        except (subprocess.CalledProcessError, OSError, ValueError) as exc:
            print(f"FAILED {slug}: {exc}", file=sys.stderr)
            failed += 1
            continue
        hub_pdf = HUB_PDF_DIR / f"{slug}.pdf"
        hub_pages = pdf_page_count(hub_pdf)
        header = [
            f"# {entry['title']}",
            "",
            f"- Slug: `{slug}`",
            f"- Label: {entry.get('label', '')}",
            f"- Week: {entry['week'] if entry['week'] is not None else 'n/a'}",
            f"- Kind: {entry['kind']}",
            f"- Original (relative to course root): `{entry['file']}`",
            f"- Study-hub PDF (relative to course root): `study-hub/public/materials/{slug}.pdf`"
            + ("" if hub_pdf.exists() else " (MISSING)"),
            f"- {unit} in original: {count if count is not None else 'n/a (plain-text conversion has no pages)'}",
            f"- Pages in study-hub PDF: {hub_pages if hub_pages is not None else 'n/a'}",
            f"- Extracted by: `06_workflow/scripts/extract_materials.py` ({method})",
            f"- OCR pages: {', '.join(map(str, ocr_pages)) if ocr_pages else 'none'}",
            "",
            "Text below is machine-extracted from the original file. Quote it only after checking the original slide or page.",
            "",
        ]
        text = "\n".join(header + body).rstrip().replace("\u00a0", " ") + "\n"
        out.write_text(text, encoding="utf-8")
        written += 1
        ocr_note = f", OCR pages: {', '.join(map(str, ocr_pages))}" if ocr_pages else ""
        print(f"wrote {out.relative_to(WS)} ({unit.lower()}: {count}, hub PDF pages: {hub_pages}{ocr_note})")
    print(f"done: {written} written, {skipped} skipped (up to date), {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
